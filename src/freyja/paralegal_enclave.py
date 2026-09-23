from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, Response
from pydantic import BaseModel, Field

paralegal_router = APIRouter(prefix="/paralegal", tags=["paralegal-enclave"])


def _default_data_dir() -> Path:
    configured = os.environ.get("FREYJA_PARALEGAL_DATA_DIR")
    if configured:
        return Path(configured).expanduser()
    vulcan_data = Path("/data/freyja/paralegal-enclave")
    if vulcan_data.parent.exists():
        return vulcan_data
    return Path.home() / ".freyja" / "paralegal-enclave"


DATA_DIR = _default_data_dir()
MATTERS_DIR = DATA_DIR / "matters"
DB_FILE = DATA_DIR / "financial_extractions.sqlite3"
EXPORT_FILE = DATA_DIR / "financial_extractions.csv"
MATTER_WORKBOOK_FILENAME = os.environ.get("FREYJA_PARALEGAL_WORKBOOK_FILENAME", "financial_extractions.xlsx")
PUBLIC_BASE_URL = os.environ.get("FREYJA_PARALEGAL_PUBLIC_BASE_URL", "").rstrip("/")
SCAN_INTERVAL_SECONDS = 10
_AUTO_SCAN_TASK: asyncio.Task[None] | None = None
SUPPORTED_DOCUMENT_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".txt"}
SUPPORTED_UPLOAD_SUFFIXES = ", ".join(sorted(SUPPORTED_DOCUMENT_SUFFIXES))

MONEY_PATTERN = re.compile(
    r"(?P<label>[A-Za-z][A-Za-z0-9 /&.,'()-]{0,80}?)?\s*"
    r"(?P<amount>\$[\s-]?\d{1,3}(?:,\d{3})*(?:\.\d{2})?|\$\s?\d+(?:\.\d{2})?)"
)


class MatterCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)


def _ensure_storage() -> None:
    MATTERS_DIR.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(DB_FILE) as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS documents (
                document_id TEXT PRIMARY KEY,
                matter_slug TEXT NOT NULL,
                filename TEXT NOT NULL,
                path TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                status TEXT NOT NULL,
                page_count INTEGER NOT NULL DEFAULT 0,
                processed_at TEXT NOT NULL,
                error TEXT
            )
            """
        )
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS financial_extractions (
                extraction_id TEXT PRIMARY KEY,
                document_id TEXT NOT NULL,
                matter_slug TEXT NOT NULL,
                filename TEXT NOT NULL,
                page_number INTEGER NOT NULL,
                field_label TEXT NOT NULL,
                amount_text TEXT NOT NULL,
                context TEXT NOT NULL,
                source_link TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(document_id) REFERENCES documents(document_id)
            )
            """
        )
        existing_columns = {
            row[1] for row in connection.execute("PRAGMA table_info(financial_extractions)")
        }
        migrations = {
            "page_text_start": "ALTER TABLE financial_extractions ADD COLUMN page_text_start INTEGER NOT NULL DEFAULT 0",
            "page_text_end": "ALTER TABLE financial_extractions ADD COLUMN page_text_end INTEGER NOT NULL DEFAULT 0",
            "text_anchor": "ALTER TABLE financial_extractions ADD COLUMN text_anchor TEXT NOT NULL DEFAULT ''",
        }
        for column, statement in migrations.items():
            if column not in existing_columns:
                connection.execute(statement)


def _matter_slug(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.strip().lower()).strip("-")
    if not slug:
        raise HTTPException(status_code=400, detail="Matter name must contain letters or numbers")
    return slug[:80]


def _matter_path(slug: str) -> Path:
    if slug != _matter_slug(slug):
        raise HTTPException(status_code=400, detail="Invalid matter folder")
    return MATTERS_DIR / slug


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _document_id(matter_slug: str, filename: str, sha256: str) -> str:
    return hashlib.sha256(f"{matter_slug}:{filename}:{sha256}".encode("utf-8")).hexdigest()[:24]


def _extract_pdf_pages(path: Path) -> tuple[int, list[tuple[int, str]]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[tuple[int, str]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            text = _ocr_pdf_page(path, index)
        pages.append((index, text))
    return len(reader.pages), pages


def _extract_image_pages(path: Path) -> tuple[int, list[tuple[int, str]]]:
    text = _ocr_image(path)
    return 1, [(1, text)]


def _extract_text_pages(path: Path) -> tuple[int, list[tuple[int, str]]]:
    return 1, [(1, path.read_text(encoding="utf-8", errors="replace"))]


def _extract_document_pages(path: Path) -> tuple[int, list[tuple[int, str]]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_pages(path)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return _extract_image_pages(path)
    if suffix == ".txt":
        return _extract_text_pages(path)
    raise ValueError(f"Unsupported document type: {suffix}")


def _ocr_pdf_page(path: Path, page_number: int) -> str:
    pdftoppm = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    if not pdftoppm or not tesseract:
        return ""
    with tempfile.TemporaryDirectory(prefix="freyja-paralegal-ocr-") as temp_dir:
        prefix = Path(temp_dir) / "page"
        render = subprocess.run(
            [
                pdftoppm,
                "-f",
                str(page_number),
                "-l",
                str(page_number),
                "-png",
                "-singlefile",
                str(path),
                str(prefix),
            ],
            capture_output=True,
            timeout=60,
            check=False,
        )
        image_path = prefix.with_suffix(".png")
        if render.returncode != 0 or not image_path.exists():
            return ""
        ocr = subprocess.run(
            [tesseract, str(image_path), "stdout"],
            capture_output=True,
            text=True,
            timeout=60,
            check=False,
        )
        if ocr.returncode != 0:
            return ""
        return ocr.stdout


def _ocr_image(path: Path) -> str:
    tesseract = shutil.which("tesseract")
    if not tesseract:
        return ""
    ocr = subprocess.run(
        [tesseract, str(path), "stdout"],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    if ocr.returncode != 0:
        return ""
    return ocr.stdout


def _financial_rows(matter_slug: str, document_id: str, filename: str, pages: list[tuple[int, str]]) -> list[dict[str, str | int]]:
    rows: list[dict[str, str | int]] = []
    for page_number, text in pages:
        compact = re.sub(r"\s+", " ", text)
        for index, match in enumerate(MONEY_PATTERN.finditer(compact)):
            label = (match.group("label") or "Financial amount").strip(" :\t\r\n")
            amount = re.sub(r"\s+", "", match.group("amount"))
            amount_start = match.start("amount")
            amount_end = match.end("amount")
            start = max(0, amount_start - 90)
            stop = min(len(compact), amount_end + 90)
            context = compact[start:stop].strip()
            anchor_start = max(0, amount_start - 24)
            anchor_stop = min(len(compact), amount_end + 24)
            text_anchor = compact[anchor_start:anchor_stop].strip()
            extraction_id = hashlib.sha256(
                f"{document_id}:{page_number}:{index}:{match.group(0)}".encode("utf-8")
            ).hexdigest()[:24]
            rows.append(
                {
                    "extraction_id": extraction_id,
                    "document_id": document_id,
                    "matter_slug": matter_slug,
                    "filename": filename,
                    "page_number": page_number,
                    "field_label": label[-80:] or "Financial amount",
                    "amount_text": amount,
                    "context": context,
                    "page_text_start": amount_start,
                    "page_text_end": amount_end,
                    "text_anchor": text_anchor,
                    "source_link": (
                        f"/paralegal/api/matters/{matter_slug}/documents/{document_id}/source"
                        f"?page={page_number}&extraction_id={extraction_id}"
                    ),
                    "created_at": datetime.now(UTC).isoformat(),
                }
            )
    return rows


def _source_link(path: str) -> str:
    if not PUBLIC_BASE_URL:
        return path
    return f"{PUBLIC_BASE_URL}{path}"


def _is_supported_document(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in SUPPORTED_DOCUMENT_SUFFIXES


def _matter_workbook_path(matter_slug: str) -> Path:
    return _matter_path(matter_slug) / MATTER_WORKBOOK_FILENAME


def _upsert_document(connection: sqlite3.Connection, payload: dict[str, Any]) -> None:
    connection.execute(
        """
        INSERT INTO documents(document_id, matter_slug, filename, path, sha256, status, page_count, processed_at, error)
        VALUES(:document_id, :matter_slug, :filename, :path, :sha256, :status, :page_count, :processed_at, :error)
        ON CONFLICT(document_id) DO UPDATE SET
            filename=excluded.filename,
            path=excluded.path,
            status=excluded.status,
            page_count=excluded.page_count,
            processed_at=excluded.processed_at,
            error=excluded.error
        """,
        payload,
    )


def process_new_documents() -> dict[str, Any]:
    _ensure_storage()
    processed = 0
    skipped = 0
    errors: list[dict[str, str]] = []
    touched_matters: set[str] = set()
    with sqlite3.connect(DB_FILE) as connection:
        known = {row[0] for row in connection.execute("SELECT document_id FROM documents")}
        for document_path in sorted(path for path in MATTERS_DIR.glob("*/*") if _is_supported_document(path)):
            matter_slug = document_path.parent.name
            touched_matters.add(matter_slug)
            sha256 = _hash_file(document_path)
            document_id = _document_id(matter_slug, document_path.name, sha256)
            if document_id in known:
                skipped += 1
                continue
            now = datetime.now(UTC).isoformat()
            try:
                page_count, pages = _extract_document_pages(document_path)
                rows = _financial_rows(matter_slug, document_id, document_path.name, pages)
                _upsert_document(
                    connection,
                    {
                        "document_id": document_id,
                        "matter_slug": matter_slug,
                        "filename": document_path.name,
                        "path": str(document_path),
                        "sha256": sha256,
                        "status": "processed",
                        "page_count": page_count,
                        "processed_at": now,
                        "error": None,
                    },
                )
                connection.executemany(
                    """
                    INSERT OR REPLACE INTO financial_extractions(
                        extraction_id, document_id, matter_slug, filename, page_number,
                        field_label, amount_text, context, page_text_start, page_text_end,
                        text_anchor, source_link, created_at
                    )
                    VALUES(
                        :extraction_id, :document_id, :matter_slug, :filename, :page_number,
                        :field_label, :amount_text, :context, :page_text_start, :page_text_end,
                        :text_anchor, :source_link, :created_at
                    )
                    """,
                    rows,
                )
                processed += 1
            except Exception as exc:
                _upsert_document(
                    connection,
                    {
                        "document_id": document_id,
                        "matter_slug": matter_slug,
                        "filename": document_path.name,
                        "path": str(document_path),
                        "sha256": sha256,
                        "status": "error",
                        "page_count": 0,
                        "processed_at": now,
                        "error": str(exc),
                    },
                )
                errors.append({"file": str(document_path), "error": str(exc)})
    write_excel_readable_export()
    for matter_slug in sorted(touched_matters):
        write_matter_workbook(matter_slug)
    return {"processed": processed, "skipped": skipped, "errors": errors}


async def _auto_scan_loop() -> None:
    while True:
        await asyncio.to_thread(process_new_documents)
        await asyncio.sleep(SCAN_INTERVAL_SECONDS)


def start_auto_scan() -> None:
    global _AUTO_SCAN_TASK
    if _AUTO_SCAN_TASK is None or _AUTO_SCAN_TASK.done():
        _AUTO_SCAN_TASK = asyncio.create_task(_auto_scan_loop())


async def stop_auto_scan() -> None:
    global _AUTO_SCAN_TASK
    if _AUTO_SCAN_TASK is None:
        return
    _AUTO_SCAN_TASK.cancel()
    try:
        await _AUTO_SCAN_TASK
    except asyncio.CancelledError:
        pass
    _AUTO_SCAN_TASK = None


def _rows() -> list[dict[str, Any]]:
    _ensure_storage()
    with sqlite3.connect(DB_FILE) as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute("SELECT * FROM financial_extractions ORDER BY created_at DESC")]


def _matter_rows(matter_slug: str) -> list[dict[str, Any]]:
    _ensure_storage()
    with sqlite3.connect(DB_FILE) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT
                f.matter_slug,
                f.filename,
                d.path AS document_path,
                f.page_number,
                f.field_label,
                f.amount_text,
                f.context,
                f.page_text_start,
                f.page_text_end,
                f.text_anchor,
                f.source_link,
                d.sha256,
                d.status,
                d.error,
                f.document_id,
                f.extraction_id,
                f.created_at
            FROM financial_extractions f
            JOIN documents d ON d.document_id = f.document_id
            WHERE f.matter_slug = ?
            ORDER BY f.created_at DESC
            """,
            (matter_slug,),
        ).fetchall()
        return [dict(row) for row in rows]


def _extraction_csv_text(rows: list[dict[str, Any]]) -> str:
    output = io.StringIO()
    fieldnames = [
        "matter_slug",
        "filename",
        "page_number",
        "field_label",
        "amount_text",
        "context",
        "page_text_start",
        "page_text_end",
        "text_anchor",
        "source_link",
        "document_id",
        "extraction_id",
        "created_at",
    ]
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue()


def write_excel_readable_export() -> Path:
    _ensure_storage()
    EXPORT_FILE.write_text(_extraction_csv_text(_rows()), encoding="utf-8")
    return EXPORT_FILE


def write_matter_workbook(matter_slug: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    _ensure_storage()
    path = _matter_workbook_path(matter_slug)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = _matter_rows(matter_slug)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Financial Extracts"
    headers = [
        "Matter",
        "Document",
        "Document Link",
        "Source Link",
        "Page",
        "Label",
        "Amount",
        "Context",
        "Text Anchor",
        "Document SHA-256",
        "Status",
        "Error",
        "Created At",
        "Document ID",
        "Extraction ID",
    ]
    worksheet.append(headers)
    for row in rows:
        worksheet.append(
            [
                row["matter_slug"],
                row["filename"],
                "Open document",
                "Open source page",
                row["page_number"],
                row["field_label"],
                row["amount_text"],
                row["context"],
                row["text_anchor"],
                row["sha256"],
                row["status"],
                row["error"] or "",
                row["created_at"],
                row["document_id"],
                row["extraction_id"],
            ]
        )
        current = worksheet.max_row
        document_uri = Path(row["document_path"]).resolve().as_uri()
        worksheet.cell(row=current, column=3).hyperlink = document_uri
        worksheet.cell(row=current, column=3).style = "Hyperlink"
        worksheet.cell(row=current, column=4).hyperlink = _source_link(row["source_link"])
        worksheet.cell(row=current, column=4).style = "Hyperlink"

    header_fill = PatternFill("solid", fgColor="1F4E5F")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    widths = [18, 28, 18, 20, 10, 24, 16, 72, 42, 18, 12, 24, 24, 28, 28]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(path)
    return path


@paralegal_router.get("", response_class=HTMLResponse)
async def paralegal_page() -> HTMLResponse:
    html_path = Path(__file__).with_name("paralegal_enclave.html")
    return HTMLResponse(html_path.read_text(encoding="utf-8"))


@paralegal_router.post("/api/matters")
async def create_matter(matter: MatterCreate) -> dict[str, Any]:
    _ensure_storage()
    slug = _matter_slug(matter.name)
    path = _matter_path(slug)
    path.mkdir(parents=True, exist_ok=True)
    workbook_path = write_matter_workbook(slug)
    return {"slug": slug, "name": matter.name.strip(), "path": str(path), "workbook": str(workbook_path)}


@paralegal_router.get("/api/matters")
async def list_matters() -> dict[str, Any]:
    _ensure_storage()
    matters = [
        {
            "slug": path.name,
            "path": str(path),
            "document_count": len([item for item in path.iterdir() if _is_supported_document(item)]),
            "workbook": str(_matter_workbook_path(path.name)),
        }
        for path in sorted(MATTERS_DIR.iterdir())
        if path.is_dir()
    ]
    return {"matters": matters}


@paralegal_router.put("/api/matters/{matter_slug}/documents/{filename}")
async def upload_document(matter_slug: str, filename: str, request: Request) -> dict[str, Any]:
    matter_path = _matter_path(matter_slug)
    if not matter_path.exists():
        raise HTTPException(status_code=404, detail="Matter folder not found")
    clean_name = Path(filename).name
    if Path(clean_name).suffix.lower() not in SUPPORTED_DOCUMENT_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"Only these document types are accepted: {SUPPORTED_UPLOAD_SUFFIXES}")
    payload = await request.body()
    if not payload:
        raise HTTPException(status_code=400, detail="PDF payload is empty")
    target = matter_path / clean_name
    target.write_bytes(payload)
    scan = process_new_documents()
    return {"saved": True, "path": str(target), "scan": scan}


@paralegal_router.post("/api/scan")
async def scan_documents() -> dict[str, Any]:
    return process_new_documents()


@paralegal_router.get("/api/status")
async def enclave_status() -> dict[str, Any]:
    start_auto_scan()
    _ensure_storage()
    scan = process_new_documents()
    with sqlite3.connect(DB_FILE) as connection:
        document_count = connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        extraction_count = connection.execute("SELECT COUNT(*) FROM financial_extractions").fetchone()[0]
    matter_workbooks = {
        path.name: str(_matter_workbook_path(path.name))
        for path in sorted(MATTERS_DIR.iterdir())
        if path.is_dir()
    }
    return {
        "auto_scan_running": _AUTO_SCAN_TASK is not None and not _AUTO_SCAN_TASK.done(),
        "scan_interval_seconds": SCAN_INTERVAL_SECONDS,
        "ocr_available": shutil.which("pdftoppm") is not None and shutil.which("tesseract") is not None,
        "matter_root": str(MATTERS_DIR),
        "database_file": str(DB_FILE),
        "excel_readable_file": str(EXPORT_FILE),
        "matter_workbook_filename": MATTER_WORKBOOK_FILENAME,
        "matter_workbooks": matter_workbooks,
        "supported_document_types": sorted(SUPPORTED_DOCUMENT_SUFFIXES),
        "document_count": document_count,
        "extraction_count": extraction_count,
        "scan": scan,
    }


@paralegal_router.get("/api/extractions")
async def list_extractions() -> dict[str, Any]:
    return {"extractions": _rows()}


@paralegal_router.get("/api/extractions.csv")
async def extractions_csv() -> Response:
    rows = _rows()
    csv_text = _extraction_csv_text(rows)
    EXPORT_FILE.write_text(csv_text, encoding="utf-8")
    return Response(
        csv_text,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="paralegal-financial-extractions.csv"'},
    )


@paralegal_router.get("/api/matters/{matter_slug}/documents/{document_id}/source")
async def source_document(matter_slug: str, document_id: str, page: int = 1, extraction_id: str | None = None) -> FileResponse:
    _ensure_storage()
    with sqlite3.connect(DB_FILE) as connection:
        row = connection.execute(
            "SELECT path, filename FROM documents WHERE matter_slug = ? AND document_id = ?",
            (matter_slug, document_id),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Document not found")
    path = Path(row[0])
    if not path.exists():
        raise HTTPException(status_code=404, detail="Source document not found on disk")
    headers = {"X-Source-Page": str(page)}
    if extraction_id:
        headers["X-Extraction-Id"] = extraction_id
    return FileResponse(path, filename=row[1], headers=headers)
